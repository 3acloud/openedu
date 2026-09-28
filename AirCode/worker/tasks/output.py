import uuid
import openpyxl
import markdown2
import qrcode
import os
import re
import random
import subprocess
import asyncio
import edge_tts

from pathlib import Path
from datetime import datetime

from gtts import gTTS
from docx import Document
from docx.oxml.ns import qn
from htmldocx import HtmlToDocx

from utilities.ask_answer import Answer
from utilities.workflow import Workflow
from worker.tasks import task




@task
def text(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)

    fields = workflow.get_node_fields(node_id)
    for field in fields:
        if field in ("输出"):
            continue
        if workflow.get_node_field_value(node_id, field) == "BranchSkip":
            workflow.update_node_field_value(node_id, "输出", "BranchSkip")
            return workflow.data

    text = workflow.get_node_field_value(node_id, "文本")
    workflow.update_node_field_value(node_id, "输出", text)
    if workflow_data["answered"]:
        Answer("\n", workflow_data["wid"], workflow_data["streamrunid"])
    else:
        workflow_data["answered"] = True
    Answer(text, workflow_data["wid"], workflow_data["streamrunid"])
    return workflow.data



@task
def document(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)

    fields = workflow.get_node_fields(node_id)
    for field in fields:
        if field in ("输出"):
            continue
        if workflow.get_node_field_value(node_id, field) == "BranchSkip":
            workflow.update_node_field_value(node_id, "输出", "BranchSkip")
            return workflow.data

    output_folder = Path("SERVERFILES")
    file_name = workflow.get_node_field_value(node_id, "文件名")
    content = workflow.get_node_field_value(node_id, "内容")
    if isinstance(content, list):
        contents = []
        for item in content:
            if item == None:
                item = ""
            item = str(item)
            item = item.encode('gbk',errors='ignore').decode('gbk').encode('utf-8').decode('utf-8')
            contents.append(item)
        contents = content
    else:
        if content == None:
            content = ""
        content = str(content)
        content = content.encode('gbk',errors='ignore').decode('gbk').encode('utf-8').decode('utf-8')
        contents = [content]
    export_type = workflow.get_node_field_value(node_id, "文档类型")

    file_guid = "{" + str(uuid.uuid4()) + "}" 

    local_file = output_folder / f"{file_guid}{export_type}"
    if local_file.exists():
        local_file = output_folder / f"{file_guid}_{datetime.now().strftime('%Y%m%d%H%M%S')}{export_type}"
    if export_type.lower().endswith((".txt", ".md", ".html", ".json", ".csv")):
        output = "\n".join(contents)
        if export_type.lower().endswith((".md")):
            if output.startswith('```markdown\n'):
                output = output[12:].lstrip()
            pos = output.rfind('\n```')
            if pos != -1:
                output = output[:pos]
        if export_type.lower().endswith((".json")):
            if output.startswith('```json\n'):
                output = output[8:].lstrip()
            pos = output.rfind('\n```')
            if pos != -1:
                output = output[:pos]
        if export_type.lower().endswith((".html")):
            pattern = r'```(.*?)\n(.*?)\n```'
            matches = re.findall(pattern, output, re.DOTALL)
            if (len(matches) > 0):
                output = ""
                for block in matches:
                    output = output + block[1]
            else:
                if output.startswith('```html\n'):
                    output = output[8:].lstrip()
                    output = re.sub(r'\n?```$', '', output)
                elif output.startswith('```'):
                    output = output[3:].lstrip()
                    output = re.sub(r'\n?```$', '', output)

            # 使用正则表达式替换
            pattern = r'^<html\s.*$'
            replacement = '<html>'
            # 逐行处理文本
            lines = output.splitlines()
            new_lines = []
            for line in lines:
                new_line = re.sub(pattern, replacement, line)
                new_lines.append(new_line)
            # 将处理后的行重新组合成文本
            output = '\n'.join(new_lines)

            if not "<html>" in output:
                output = "<html>\n" + output + "\n</html>"
    
            
        with open(local_file, "wb") as txt_file:
            txt_file.write(output.encode("utf-8"))
    elif export_type.lower().endswith(".docx"):
        content_str = "\n".join(contents)
        html_content = markdown2.markdown(content_str)
        new_parser = HtmlToDocx()
        document = Document()
        new_parser.add_html_to_document(html_content, document)
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.name = "微软雅黑"
                r = run._element.rPr.rFonts
                r.set(qn("w:eastAsia"), "微软雅黑")

        document.save(local_file)
    elif export_type.lower().endswith(".xlsx"):
        content_str = "\n".join(contents)
        lines = content_str.split("\n")
        wb = openpyxl.Workbook()
        ws = wb.active
        for line in lines:
            ws.append(line.split(","))
        wb.save(local_file)
    elif export_type.lower().endswith((".wav",".ogg")):
        content_str = "\n".join(contents)
        tts = gTTS(text=content_str, lang='zh-CN')
        tts.save(local_file)
    elif export_type.lower().endswith((".jpg",".jpeg",".png",".gif",".bmp",".tif")):
        content_str = "\n".join(contents)
        img = qrcode.make(content_str)
        img.save(local_file)
    else:  
        with open(local_file, "wb") as txt_file:
            txt_file.write(("\n".join(contents)).encode("utf-8"))

    basename = os.path.basename(local_file)
    workflow.data["tempfiles"].append("SERVERFILES/" + basename)
    workflow.update_node_field_value(node_id, "输出", basename)
    if workflow_data["answered"]:
        Answer("\n", workflow_data["wid"], workflow_data["streamrunid"])
    else:
        workflow_data["answered"] = True
    if file_name == None or file_name == "":
        Answer("File://" + file_guid + export_type, workflow_data["wid"], workflow_data["streamrunid"])
    else:
        Answer("File://" + file_name + export_type + file_guid + export_type, workflow_data["wid"], workflow_data["streamrunid"])
    return workflow.data



def text_to_speech(text, voice="zh-CN-YunxiNeural", output_file="output.mp3"):
    """同步版本的文本转语音函数"""
    async def _internal():
        communicate = edge_tts.Communicate(text, voice)
        await communicate.save(output_file)
    
    asyncio.run(_internal())


@task
def voice(
    workflow_data: dict,
    node_id: str,
):
    workflow = Workflow(workflow_data)

    fields = workflow.get_node_fields(node_id)
    for field in fields:
        if field in ("输出"):
            continue
        if workflow.get_node_field_value(node_id, field) == "BranchSkip":
            workflow.update_node_field_value(node_id, "输出", "BranchSkip")
            return workflow.data

    output_folder = Path("SERVERFILES")
    file_name = workflow.get_node_field_value(node_id, "文件名")
    if file_name == None or file_name == "":
        file_name = "音频"
    content = workflow.get_node_field_value(node_id, "内容")
    if isinstance(content, list):
        contents = []
        for item in content:
            if item == None:
                item = ""
            item = str(item)
            item = item.encode('gbk',errors='ignore').decode('gbk').encode('utf-8').decode('utf-8')
            contents.append(item)
        contents = content
    else:
        if content == None:
            content = ""
        content = str(content)
        content = content.encode('gbk',errors='ignore').decode('gbk').encode('utf-8').decode('utf-8')
        contents = [content]
    export_type = ".mp3"

    file_guid = "{" + str(uuid.uuid4()) + "}" 

    local_file = output_folder / f"{file_guid}{export_type}"
    if local_file.exists():
        local_file = output_folder / f"{file_guid}_{datetime.now().strftime('%Y%m%d%H%M%S')}{export_type}"

    content_str = "\n".join(contents)
    voice = workflow.get_node_field_value(node_id, "音色")
    if voice == None or voice == "":
        voice = random.choice("eve;ara;rex;sal;leo".split(";"))
        voice = random.choice("zh-CN-YunxiNeural;zh-CN-YunyangNeural;zh-CN-YunjianNeural;zh-CN-XiaoxiaoNeural;zh-CN-XiaoyiNeural;zh-CN-XiaomengNeural".split(";"))
    if voice not in "zh-CN-YunxiNeural;zh-CN-YunyangNeural;zh-CN-YunjianNeural;zh-CN-XiaoxiaoNeural;zh-CN-XiaoyiNeural;zh-CN-XiaomengNeural":
        voice = "zh-CN-YunxiNeural"
    if voice in "zh-CN-YunxiNeural;zh-CN-YunyangNeural;zh-CN-YunjianNeural;zh-CN-XiaoxiaoNeural;zh-CN-XiaoyiNeural;zh-CN-XiaomengNeural":
        try:
            text_to_speech(content_str, voice, local_file)
        except Exception:
            pass
        
    if local_file.exists():
        basename = os.path.basename(local_file)
        workflow.data["tempfiles"].append("SERVERFILES/" + basename)
        workflow.update_node_field_value(node_id, "输出", basename)
    if workflow_data["answered"]:
        Answer("\n", workflow_data["wid"], workflow_data["streamrunid"])
    else:
        workflow_data["answered"] = True
    if local_file.exists():
        Answer("File://" + file_name + export_type + file_guid + export_type, workflow_data["wid"], workflow_data["streamrunid"])
    return workflow.data


